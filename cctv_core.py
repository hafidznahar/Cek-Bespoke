"""
Logika inti CCTV Quality Checker (dipindahkan dari cctv_checker.py).
Threshold & algoritma SAMA dengan script asli; yang berubah hanya cara input/output
(tidak lagi memakai path D:\\ / argparse, tetapi folder sementara hasil upload).
"""
import hashlib
import io
import re
import shutil
from pathlib import Path

import cv2
import numpy as np
from PIL import Image

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

INPUT_EXTENSIONS = (".pptx", ".pdf")
FOLDER_HASIL_NAME = "Hasil"

# =========================== CONFIG (TUNED FOR REAL CCTV) ===========================

# --- 1. RESOLUSI MINIMAL ---
MIN_WIDTH_PX = 320
MIN_HEIGHT_PX = 200
MIN_TOTAL_PIXELS = 80000

# --- 2. KETAJAMAN & BLUR (Gaussian-Filtered Laplacian) ---
FILTERED_BLUR_THRESHOLD = 28.0
RAW_BLUR_THRESHOLD = 45.0

# --- 3. COMPRESSION NOISE / GAMBAR PECAH ---
MAX_NOISE_RATIO = 12.0

# --- 4. KONTRAS & WASHED-OUT (PUCAT) ---
MIN_STD_CONTRAST = 22.0
MIN_DYNAMIC_RANGE = 65.0

# --- 5. EXPOSURE (SILAU / GELAP GULITA) ---
MAX_OVEREXPOSED_PCT = 35.0
MAX_UNDEREXPOSED_PCT = 55.0

# --- 6. FISHEYE ---
FISHEYE_DARK_THRESHOLD = 35
FISHEYE_MIN_CIRCULARITY = 0.85
FISHEYE_MAX_RADIUS_RATIO = 1.05
FISHEYE_MIN_BLACK_FRACTION = 0.15

# --- 7. EXCEL & LOGO FILTER ---
THUMB_WIDTH = 200
LOGO_TRANSPARENT_RATIO = 0.15
LOGO_MAX_UNIQUE_COLORS = 1500
EXTRA_EXCLUDE_HASHES = set()


# ====================================================================================


def log(msg):
    print(msg, flush=True)


# ------------------------------- PENCARIAN FILE -------------------------------

def find_input_files(input_dir: Path):
    """Cari semua .pptx dan .pdf di folder input dan semua subfoldernya."""
    found = []
    for p in input_dir.rglob("*"):
        if not p.is_file():
            continue
        if p.name.startswith("~$"):
            continue
        if p.suffix.lower() not in INPUT_EXTENSIONS:
            continue
        rel_parts = p.relative_to(input_dir).parts[:-1]
        # lewati folder hasil, folder sementara, dan folder tersembunyi
        if any(part == FOLDER_HASIL_NAME or part.startswith(("_temp", ".")) for part in rel_parts):
            continue
        found.append(p)
    return sorted(found, key=lambda x: str(x).lower())


# ------------------------------- NAMA CABANG -------------------------------

BRANCH_PATTERN = re.compile(r"\bKCP?\b[\s._-]*[A-Za-z][A-Za-z .]*", re.IGNORECASE)


def sanitize_filename(s: str):
    return re.sub(r"[^A-Za-z0-9_-]+", "_", s).strip("_")


def make_image_name(row):
    parts = [f"{row['doc_idx']:03d}", row["branch"], f"s{row['slide']}", f"img{row['img_index']}"]
    if row.get("label"):
        parts.append(row["label"])
    name = sanitize_filename("_".join(str(p) for p in parts))
    return name[:120] + ".jpg"


def branch_name_from_filename(path: Path):
    stem = path.stem
    stem = re.sub(r"[_\-]+", " ", stem)
    m = BRANCH_PATTERN.search(stem)
    if m:
        return re.sub(r"\s+", " ", m.group(0)).strip().upper()
    return stem.strip()


def branch_name_from_text(text, fallback):
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        m = BRANCH_PATTERN.search(line)
        if m:
            return re.sub(r"\s+", " ", m.group(0)).strip().upper()
    return fallback


# ------------------------------- BACA GAMBAR DARI PPTX / PDF -------------------------------

def iter_shapes(shapes):
    """Telusuri semua shape, termasuk yang ada di dalam group."""
    for shape in shapes:
        yield shape
        try:
            if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
                yield from iter_shapes(shape.shapes)
        except Exception:
            pass


def iter_pptx_images(path: Path, need_branch: bool):
    """Yield (nomor_slide, nama_cabang, blob_gambar)."""
    if Presentation is None:
        raise RuntimeError("python-pptx belum terpasang (pip install python-pptx)")
    fallback = branch_name_from_filename(path)
    prs = Presentation(str(path))
    for slide_idx, slide in enumerate(prs.slides, start=1):
        branch = fallback
        if need_branch:
            texts = []
            for shape in iter_shapes(slide.shapes):
                if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                    texts.append(shape.text_frame.text)
            branch = branch_name_from_text("\n".join(texts), fallback)
        for shape in iter_shapes(slide.shapes):
            try:
                blob = shape.image.blob
            except Exception:
                continue
            yield slide_idx, branch, blob


def iter_pdf_images(path: Path, need_branch: bool):
    """Yield (nomor_halaman, nama_cabang, blob_gambar)."""
    if fitz is None:
        raise RuntimeError("PyMuPDF belum terpasang (pip install pymupdf)")
    fallback = branch_name_from_filename(path)
    doc = fitz.open(str(path))
    seen_xrefs = set()
    try:
        for page_no, page in enumerate(doc, start=1):
            branch = fallback
            if need_branch:
                branch = branch_name_from_text(page.get_text(), fallback)
            for img in page.get_images(full=True):
                xref = img[0]
                if xref in seen_xrefs:
                    continue
                seen_xrefs.add(xref)
                try:
                    info = doc.extract_image(xref)
                except Exception:
                    continue
                if info and info.get("image"):
                    yield page_no, branch, info["image"]
    finally:
        doc.close()


def iter_document_images(path: Path, need_branch: bool):
    if path.suffix.lower() == ".pptx":
        yield from iter_pptx_images(path, need_branch)
    else:
        yield from iter_pdf_images(path, need_branch)


def collect_image_hashes(paths):
    """Pass 1: hash semua gambar embedded di semua file (untuk membuang logo/template)."""
    hash_counts = {}
    for p in paths:
        try:
            for _, _, blob in iter_document_images(p, need_branch=False):
                h = hashlib.md5(blob).hexdigest()
                hash_counts[h] = hash_counts.get(h, 0) + 1
        except Exception as e:
            log(f"  [WARN] Gagal baca {p.name}: {e}")
    return hash_counts


# ------------------------------- FILTER & ANALISIS -------------------------------

def pil_to_cv2(pil_img: Image.Image):
    if pil_img.mode != "RGB":
        pil_img = pil_img.convert("RGB")
    arr = np.array(pil_img)
    return cv2.cvtColor(arr, cv2.COLOR_RGB2BGR)


def is_blank_or_graphic(pil_img: Image.Image):
    """Mendeteksi & membuang elemen UI, tombol, shape, background slide, atau gambar tidak relevan."""
    gray = np.array(pil_img.convert("L"))
    std_dev = float(gray.std())

    # 1. Gambar polos / 1 warna solid
    if std_dev < 12.0:
        return True

    # 2. Didominasi background putih / UI (>65% piksel bernilai >235)
    white_ratio = float((gray > 235).mean())
    if white_ratio > 0.65:
        return True

    # 3. Didominasi layar gelap / frame hitam mati (>80%)
    black_ratio = float((gray < 20).mean())
    if black_ratio > 0.80:
        return True

    return False


def assess_camera_quality(cv_img):
    """Evaluasi Kualitas Gambar Kamera CCTV."""
    h, w, _ = cv_img.shape
    gray = cv2.cvtColor(cv_img, cv2.COLOR_BGR2GRAY)

    # Potong margin hitam luar
    thresh = cv2.threshold(gray, 20, 255, cv2.THRESH_BINARY)[1]
    contours, _ = cv2.findContours(thresh, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    if contours:
        c_max = max(contours, key=cv2.contourArea)
        x, y, bw, bh = cv2.boundingRect(c_max)
        if bw > 60 and bh > 60:
            roi_gray = gray[y:y + bh, x:x + bw]
            roi_w, roi_h = bw, bh
        else:
            roi_gray, roi_w, roi_h = gray, w, h
    else:
        roi_gray, roi_w, roi_h = gray, w, h

    # Margin 8% mengabaikan timestamp/logo teks di pojok
    margin_y = int(roi_h * 0.08)
    margin_x = int(roi_w * 0.08)
    core_gray = roi_gray[margin_y:max(margin_y + 1, roi_h - margin_y), margin_x:max(margin_x + 1, roi_w - margin_x)]
    if core_gray.size == 0:
        core_gray = roi_gray

    # 1. Check Resolusi
    total_pixels = roi_w * roi_h
    is_low_res = (roi_w < MIN_WIDTH_PX or roi_h < MIN_HEIGHT_PX or total_pixels < MIN_TOTAL_PIXELS)

    # 2. Check Ketajaman & Noise
    raw_blur = cv2.Laplacian(core_gray, cv2.CV_64F).var()
    blurred_core = cv2.GaussianBlur(core_gray, (3, 3), 0)
    filtered_blur = cv2.Laplacian(blurred_core, cv2.CV_64F).var()
    noise_ratio = raw_blur / (filtered_blur + 1e-5)

    is_blur = (filtered_blur < FILTERED_BLUR_THRESHOLD) or (raw_blur < RAW_BLUR_THRESHOLD)
    is_pixelated_noise = (noise_ratio > MAX_NOISE_RATIO) and (filtered_blur < 100.0 or is_low_res)

    # 3. Check Kontras / Pucat
    std_contrast = float(np.std(core_gray))
    p5, p95 = np.percentile(core_gray, [5, 95])
    dynamic_range = float(p95 - p5)
    is_washed_out = (std_contrast < MIN_STD_CONTRAST) or (dynamic_range < MIN_DYNAMIC_RANGE)

    # 4. Check Exposure
    overexposed_pct = float((core_gray > 245).mean() * 100)
    underexposed_pct = float((core_gray < 15).mean() * 100)
    is_bad_exposure = (overexposed_pct > MAX_OVEREXPOSED_PCT) or (underexposed_pct > MAX_UNDEREXPOSED_PCT)

    issues = []
    if is_low_res:
        issues.append(f"Resolusi kecil ({roi_w}x{roi_h}px)")
    if is_blur:
        issues.append(f"Gambar burem (Score: {filtered_blur:.1f})")
    if is_pixelated_noise:
        issues.append(f"Gambar pecah/noise (Ratio: {noise_ratio:.1f})")
    if is_washed_out:
        issues.append(f"Gambar pucat/kontras rendah ({std_contrast:.1f})")
    if is_bad_exposure:
        if overexposed_pct > MAX_OVEREXPOSED_PCT:
            issues.append(f"Overexposed/Silau ({overexposed_pct:.1f}%)")
        if underexposed_pct > MAX_UNDEREXPOSED_PCT:
            issues.append(f"Terlalu gelap ({underexposed_pct:.1f}%)")

    is_ok = len(issues) == 0

    detail = {
        "roi_w": roi_w,
        "roi_h": roi_h,
        "raw_blur": round(raw_blur, 2),
        "filtered_blur": round(filtered_blur, 2),
        "noise_ratio": round(noise_ratio, 2),
        "std_contrast": round(std_contrast, 2),
        "dynamic_range": round(dynamic_range, 2),
        "overexposed_pct": round(overexposed_pct, 1),
        "underexposed_pct": round(underexposed_pct, 1),
    }

    return is_ok, issues, detail


def fisheye_check(gray):
    """Deteksi Fisheye Ketat."""
    h, w = gray.shape
    content_mask = (gray > FISHEYE_DARK_THRESHOLD).astype(np.uint8) * 255
    content_mask = cv2.morphologyEx(content_mask, cv2.MORPH_CLOSE, np.ones((5, 5), np.uint8))

    contours, _ = cv2.findContours(content_mask, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    black_fraction = 1.0 - (content_mask > 0).mean()

    if not contours:
        return 0.0, False, {"circularity": 0.0, "radius_ratio": 0.0, "black_fraction": round(black_fraction, 3)}

    c = max(contours, key=cv2.contourArea)
    area = cv2.contourArea(c)
    (_, _), radius = cv2.minEnclosingCircle(c)
    circle_area = np.pi * radius * radius if radius > 0 else 1.0
    circularity = min(1.0, area / circle_area)
    radius_ratio = radius / (min(h, w) / 2.0)

    is_fisheye = (
            circularity >= FISHEYE_MIN_CIRCULARITY
            and radius_ratio <= FISHEYE_MAX_RADIUS_RATIO
            and black_fraction >= FISHEYE_MIN_BLACK_FRACTION
    )

    detail = {
        "circularity": round(circularity, 3),
        "radius_ratio": round(radius_ratio, 3),
        "black_fraction": round(black_fraction, 3),
    }
    return circularity, is_fisheye, detail


def is_logo_like(pil_img: Image.Image):
    im = pil_img.convert("RGBA")
    small = im.resize((150, 150))
    arr = np.array(small)
    alpha = arr[:, :, 3]
    transparent_ratio = (alpha < 10).mean()
    if transparent_ratio > LOGO_TRANSPARENT_RATIO:
        return True, transparent_ratio, None

    flat = arr.reshape(-1, 4)
    unique_colors = len(np.unique(flat, axis=0))
    if unique_colors < LOGO_MAX_UNIQUE_COLORS:
        return True, transparent_ratio, unique_colors

    return False, transparent_ratio, unique_colors


# ------------------------------- SPLIT GRID DASHBOARD -------------------------------

def _content_and_gutter_segments(frac, gutter_thr, min_gap_merge, min_len):
    gutter = frac > gutter_thr
    raw_segs = []
    start = None
    for i, g in enumerate(gutter):
        if g and start is None:
            start = i
        if not g and start is not None:
            raw_segs.append([start, i])
            start = None
    if start is not None:
        raw_segs.append([start, len(gutter)])

    merged = []
    for s in raw_segs:
        if merged and s[0] - merged[-1][1] <= min_gap_merge:
            merged[-1][1] = s[1]
        else:
            merged.append(s)

    content = []
    prev_end = 0
    for s in merged:
        if s[0] - prev_end >= min_len:
            content.append((prev_end, s[0]))
        prev_end = s[1]
    if len(frac) - prev_end >= min_len:
        content.append((prev_end, len(frac)))
    return content, merged


def _smooth(x, k=11):
    if k <= 1:
        return x
    kernel = np.ones(k) / k
    return np.convolve(x, kernel, mode="same")


def ocr_label(crop_img):
    try:
        import pytesseract
        text = pytesseract.image_to_string(crop_img, config="--psm 6")
    except Exception:
        return None
    words = re.findall(r"[A-Za-z]{3,}", text)
    if not words:
        return None
    return words[0].capitalize()


def split_dashboard_grid(pil_img, min_tile=120):
    arr = np.array(pil_img.convert("RGB"))
    gray = arr.mean(axis=2)
    h, w = gray.shape

    col_frac = (gray > 225).astype(float).mean(axis=0)
    row_frac = _smooth((gray > 225).astype(float).mean(axis=1), 11)

    col_content, _ = _content_and_gutter_segments(col_frac, 0.90, 8, min_tile)
    row_content, row_gutters = _content_and_gutter_segments(row_frac, 0.85, 20, min_tile)

    if len(col_content) * len(row_content) <= 1:
        return None

    tiles = []
    for ri, (ry0, ry1) in enumerate(row_content):
        label_band = None
        for g in row_gutters:
            if g[0] >= ry1 - 2:
                label_band = (g[0], min(h, g[0] + 90))
                break
        for ci, (cx0, cx1) in enumerate(col_content):
            tile_img = pil_img.crop((cx0, ry0, cx1, ry1))
            if is_blank_or_graphic(tile_img):
                continue
            label = None
            if label_band:
                label_crop = pil_img.crop((cx0, label_band[0], cx1, label_band[1]))
                label = ocr_label(label_crop)
            tiles.append({
                "tile_img": tile_img,
                "row": ri + 1,
                "col": ci + 1,
                "label": label or f"Cam_R{ri + 1}C{ci + 1}",
            })
    return tiles if tiles else None


# ------------------------------- EKSTRAK KANDIDAT SCREENSHOT -------------------------------

def extract_candidate_images(path: Path, hash_counts, doc_idx, folder_label, logo_threshold=2):
    results = []
    img_idx = 0

    for page_idx, branch, blob in iter_document_images(path, need_branch=True):
        h = hashlib.md5(blob).hexdigest()
        if hash_counts.get(h, 0) >= logo_threshold or h in EXTRA_EXCLUDE_HASHES:
            continue

        try:
            pil_img = Image.open(io.BytesIO(blob))
            pil_img.load()
        except Exception:
            continue

        w, h_px = pil_img.size
        if w < 80 or h_px < 80:
            continue

        # Filter Logo & Shape/Graphic UI
        logo_like, _, _ = is_logo_like(pil_img)
        if logo_like or is_blank_or_graphic(pil_img):
            continue

        img_idx += 1
        tiles = split_dashboard_grid(pil_img)

        base = {
            "doc_idx": doc_idx,
            "folder": folder_label,
            "branch": branch,
            "file": path.name,
            "slide": page_idx,
            "img_index": img_idx,
        }

        if tiles is None:
            results.append({**base, "label": None, "pil_img": pil_img, "width": w, "height": h_px})
        else:
            for t in tiles:
                tw, th = t["tile_img"].size
                results.append({**base, "label": t["label"], "pil_img": t["tile_img"], "width": tw, "height": th})
    return results


# ------------------------------- LAPORAN EXCEL -------------------------------

def build_report(all_rows, output_path, thumb_dir):
    wb = Workbook()

    # ---------------- Sheet 1: Hasil QC ----------------
    ws = wb.active
    ws.title = "Hasil QC"

    headers = ["No", "Folder", "Cabang", "File Sumber", "Slide/Hal", "Kamera/Lokasi", "Nama Gambar",
               "Screenshot", "Kualitas Kamera", "Fisheye", "Status Akhir", "Catatan"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ok_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    na_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    ok_font = Font(color="006100", bold=True)
    na_font = Font(color="9C0006", bold=True)

    row_num = 2
    for i, row in enumerate(all_rows, start=1):
        img_name = make_image_name(row)
        thumb_path = thumb_dir / img_name

        quality_ok = row["quality_ok"]
        fisheye_ok = row["fisheye_ok"]
        overall_ok = quality_ok and fisheye_ok

        notes = list(row["quality_issues"])
        if not fisheye_ok:
            notes.append(f"Terdeteksi Fisheye (Score: {row['fisheye_score']:.2f})")

        note_text = " | ".join(notes) if notes else "Kualitas Kamera Normal / OK"

        ws.cell(row=row_num, column=1, value=i)
        ws.cell(row=row_num, column=2, value=row["folder"])
        ws.cell(row=row_num, column=3, value=row["branch"])
        ws.cell(row=row_num, column=4, value=row["file"])
        ws.cell(row=row_num, column=5, value=row["slide"])
        ws.cell(row=row_num, column=6, value=row.get("label") or "-")
        ws.cell(row=row_num, column=7, value=img_name)

        if thumb_path.exists():
            xl_img = XLImage(str(thumb_path))
            xl_img.width = THUMB_WIDTH
            xl_img.height = int(THUMB_WIDTH * row["height"] / row["width"])
            ws.row_dimensions[row_num].height = max(55, xl_img.height * 0.75)
            ws.add_image(xl_img, f"H{row_num}")

        c_quality = ws.cell(row=row_num, column=9, value=("OK" if quality_ok else "N/A"))
        c_fisheye = ws.cell(row=row_num, column=10, value=("OK" if fisheye_ok else "N/A"))
        c_overall = ws.cell(row=row_num, column=11, value=("OK" if overall_ok else "N/A"))
        ws.cell(row=row_num, column=12, value=note_text)

        for c in (c_quality, c_fisheye, c_overall):
            if c.value == "OK":
                c.fill = ok_fill
                c.font = ok_font
            else:
                c.fill = na_fill
                c.font = na_font
            c.alignment = Alignment(horizontal="center", vertical="center")

        ws.cell(row=row_num, column=1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_num, column=5).alignment = Alignment(horizontal="center", vertical="center")
        for col in (2, 3, 4, 6, 7):
            ws.cell(row=row_num, column=col).alignment = Alignment(vertical="center", wrap_text=True)
        ws.cell(row=row_num, column=12).alignment = Alignment(vertical="center", wrap_text=True)

        row_num += 1

    widths = {"A": 5, "B": 16, "C": 22, "D": 32, "E": 9, "F": 16, "G": 30, "H": THUMB_WIDTH / 6.5,
              "I": 15, "J": 12, "K": 13, "L": 45}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{max(row_num - 1, 1)}"

    # ---------------- Sheet 2: Detail Skor (untuk tuning) ----------------
    ws2 = wb.create_sheet("Detail Skor (Tuning)")
    headers2 = [
        "Folder", "Cabang", "File Sumber", "Slide/Hal", "Kamera/Lokasi", "Nama Gambar", "Width", "Height",
        "Filtered Blur", "Blur Thr", "Noise Ratio", "Noise Max", "Contrast (Std)", "Contrast Min",
        "Overexposed %", "Underexposed %", "Circularity", "Radius Ratio", "Black Fraction"
    ]
    ws2.append(headers2)
    for col in range(1, len(headers2) + 1):
        cell = ws2.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font

    for row in all_rows:
        img_name = make_image_name(row)
        q = row["quality_detail"]
        f_d = row["fisheye_detail"]
        ws2.append([
            row["folder"], row["branch"], row["file"], row["slide"], row.get("label") or "-", img_name,
            q["roi_w"], q["roi_h"],
            q["filtered_blur"], FILTERED_BLUR_THRESHOLD,
            q["noise_ratio"], MAX_NOISE_RATIO,
            q["std_contrast"], MIN_STD_CONTRAST,
            q["overexposed_pct"], q["underexposed_pct"],
            f_d["circularity"], f_d["radius_ratio"], f_d["black_fraction"]
        ])

    for i, col in enumerate(headers2, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = max(13, len(col) + 2)

    wb.save(output_path)


# ------------------------------- MAIN (versi web) -------------------------------

def run_cctv_check(input_dir: Path, work_dir: Path, progress=None, max_preview=40):
    """
    Return dict: excel_bytes, n_docs, n_total, n_ok, n_na, summary (list dict), preview_na, warnings
    """
    warnings = []
    if Presentation is None:
        warnings.append("python-pptx belum terpasang, file .pptx tidak bisa dibaca.")
    if fitz is None:
        warnings.append("PyMuPDF belum terpasang, file .pdf tidak bisa dibaca.")

    files = find_input_files(input_dir)
    if not files:
        raise ValueError("Tidak ada file .pptx / .pdf yang ditemukan di file yang diupload.")

    out_xlsx = work_dir / "hasil_qc_cctv.xlsx"
    thumb_dir = work_dir / "_temp_thumbnails"
    if thumb_dir.exists():
        shutil.rmtree(thumb_dir)
    thumb_dir.mkdir(parents=True)

    if progress:
        progress(0.03, "Membaca semua gambar untuk filter logo/template...")
    hash_counts = collect_image_hashes(files)

    all_rows = []
    try:
        for doc_idx, p in enumerate(files, start=1):
            rel = p.relative_to(input_dir)
            folder_label = str(rel.parent) if str(rel.parent) != "." else "-"
            if progress:
                progress(0.10 + 0.70 * (doc_idx - 1) / len(files), f"Memproses {doc_idx}/{len(files)}: {rel}")

            try:
                candidates = extract_candidate_images(p, hash_counts, doc_idx, folder_label)
            except Exception as e:
                warnings.append(f"Gagal memproses {p.name}: {e}")
                continue

            for row in candidates:
                cv_img = pil_to_cv2(row["pil_img"])
                gray = cv2.cvtColor(cv_img, cv2.COLOR_BGR2GRAY)

                quality_ok, quality_issues, quality_detail = assess_camera_quality(cv_img)
                fisheye_score, is_fisheye, fisheye_detail = fisheye_check(gray)

                thumb_path = thumb_dir / make_image_name(row)
                thumb_img = row["pil_img"].convert("RGB")
                thumb_img.thumbnail((500, 500))
                thumb_img.save(thumb_path, "JPEG", quality=85)

                all_rows.append({
                    **row,
                    "quality_ok": quality_ok,
                    "quality_issues": quality_issues,
                    "quality_detail": quality_detail,
                    "fisheye_score": fisheye_score,
                    "fisheye_ok": not is_fisheye,
                    "fisheye_detail": fisheye_detail,
                })

        if not all_rows:
            raise ValueError("Tidak ada screenshot CCTV yang terdeteksi. Cek lagi isi dokumennya.")

        if progress:
            progress(0.88, "Membuat laporan Excel...")
        build_report(all_rows, out_xlsx, thumb_dir)
        excel_bytes = out_xlsx.read_bytes()

        summary, preview_na = [], []
        for i, r in enumerate(all_rows, start=1):
            overall_ok = r["quality_ok"] and r["fisheye_ok"]
            notes = list(r["quality_issues"])
            if not r["fisheye_ok"]:
                notes.append(f"Terdeteksi Fisheye (Score: {r['fisheye_score']:.2f})")
            summary.append({
                "No": i,
                "Folder": r["folder"],
                "Cabang": r["branch"],
                "File Sumber": r["file"],
                "Slide/Hal": r["slide"],
                "Kamera/Lokasi": r.get("label") or "-",
                "Kualitas Kamera": "OK" if r["quality_ok"] else "N/A",
                "Fisheye": "OK" if r["fisheye_ok"] else "N/A",
                "Status Akhir": "OK" if overall_ok else "N/A",
                "Catatan": " | ".join(notes) if notes else "Kualitas Kamera Normal / OK",
            })
            if not overall_ok and len(preview_na) < max_preview:
                tp = thumb_dir / make_image_name(r)
                if tp.exists():
                    preview_na.append({
                        "caption": f"{r['branch']} - {r['file']} (hal {r['slide']}, {r.get('label') or '-'})",
                        "note": summary[-1]["Catatan"],
                        "bytes": tp.read_bytes(),
                    })
    finally:
        shutil.rmtree(thumb_dir, ignore_errors=True)

    n_ok = sum(1 for s in summary if s["Status Akhir"] == "OK")
    if progress:
        progress(1.0, "Selesai")
    return {
        "excel_bytes": excel_bytes,
        "n_docs": len(files),
        "n_total": len(summary),
        "n_ok": n_ok,
        "n_na": len(summary) - n_ok,
        "summary": summary,
        "preview_na": preview_na,
        "warnings": warnings,
    }
