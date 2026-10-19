"""
Logika inti BAST Checker (dipindahkan dari bast_checker_pdf.py).
Threshold & algoritma SAMA dengan script asli; yang berubah hanya cara input/output
(tidak lagi memakai path D:\\ tetapi folder sementara hasil upload).
"""
import io
import os
import re
import shutil
import zipfile
from collections import defaultdict
from pathlib import Path

import imagehash
from PIL import Image, ImageStat

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

INPUT_EXTENSIONS = (".pdf", ".docx")

HASH_DISTANCE_THRESHOLD = 5
MAX_DOC_OCCURRENCES = 2  # foto yang muncul di > N dokumen dianggap header/kop/template

# filter foto
MIN_PHOTO_SIZE = 150
MIN_ASPECT_RATIO = 0.3
MAX_ASPECT_RATIO = 3.0
MIN_SATURATION = 5.0
MAX_BACKGROUND_RATIO = 0.75
MIN_STD_DEV = 8.0

# excel
THUMB_MAX_W = 130
THUMB_MAX_H = 95
ROW_HEIGHT_PT = 80
COL_WIDTH_SCREENSHOT = 20

FILL_DUPLIKAT = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
FONT_DUPLIKAT = Font(color="9C0006", bold=True)
FILL_OK = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
FONT_OK = Font(color="006100", bold=True)
FILL_HEADER = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
FONT_HEADER = Font(color="FFFFFF", bold=True)
THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)


def pdf_support_available():
    return fitz is not None


def slugify(text):
    base = os.path.splitext(text)[0]
    base = re.sub(r"[\\/]+", "_", base)
    return re.sub(r"(?u)[^-\w. ]", "", base).strip()


def collect_input_files(input_dir: Path):
    """Cari semua PDF/DOCX di folder input dan subfoldernya -> [(full, rel)]."""
    found = []
    for root, dirs, files in os.walk(input_dir):
        dirs[:] = [d for d in dirs if not d.startswith(".") and not d.startswith("__")]
        for f in files:
            if f.startswith("~$"):
                continue
            if f.lower().endswith(INPUT_EXTENSIONS):
                full = os.path.join(root, f)
                rel = os.path.relpath(full, input_dir)
                found.append((full, rel))
    return sorted(found, key=lambda x: x[1].lower())


def get_header_footer_media(z):
    hf_media = set()
    try:
        rels_files = [f for f in z.namelist() if "word/_rels/header" in f or "word/_rels/footer" in f]
        for rels in rels_files:
            content = z.read(rels).decode("utf-8", errors="ignore")
            for m in re.findall(r'Target="([^"]*media/[^"]*)"', content):
                hf_media.add(os.path.basename(m))
    except Exception:
        pass
    return hf_media


def extract_photos_from_docx(docx_path, temp_folder, prefix, warnings):
    extracted = []
    try:
        with zipfile.ZipFile(docx_path, "r") as z:
            hf_media_names = get_header_footer_media(z)
            media_files = [f for f in z.namelist() if f.startswith("word/media/")]
            for m in media_files:
                fname = os.path.basename(m)
                if fname in hf_media_names:
                    continue
                out_path = os.path.join(temp_folder, f"{prefix}__{fname}")
                with z.open(m) as src, open(out_path, "wb") as dst:
                    dst.write(src.read())
                extracted.append((out_path, fname))
    except Exception as e:
        warnings.append(f"Gagal ekstrak {os.path.basename(docx_path)}: {e}")
    return extracted


def extract_photos_from_pdf(pdf_path, temp_folder, prefix, warnings):
    extracted = []
    if fitz is None:
        warnings.append(f"PyMuPDF belum terpasang, PDF dilewati: {os.path.basename(pdf_path)}")
        return extracted
    try:
        doc = fitz.open(pdf_path)
        seen_xrefs = set()
        for page_no, page in enumerate(doc, start=1):
            for img in page.get_images(full=True):
                xref = img[0]
                if xref in seen_xrefs:
                    continue
                seen_xrefs.add(xref)
                try:
                    info = doc.extract_image(xref)
                except Exception:
                    continue
                if not info or not info.get("image"):
                    continue
                ext = info.get("ext", "png")
                label = f"hal{page_no}_img{xref}.{ext}"
                out_path = os.path.join(temp_folder, f"{prefix}__{label}")
                with open(out_path, "wb") as f:
                    f.write(info["image"])
                extracted.append((out_path, label))
        doc.close()
    except Exception as e:
        warnings.append(f"Gagal ekstrak {os.path.basename(pdf_path)}: {e}")
    return extracted


def get_saturation_mean(img):
    return ImageStat.Stat(img.convert("HSV")).mean[1]


def dominant_background_ratio(img):
    gray = img.convert("L").resize((100, 100))
    pixels = list(gray.getdata())
    counts = {}
    for p in pixels:
        b = p // 16
        counts[b] = counts.get(b, 0) + 1
    return max(counts.values()) / len(pixels)


def is_real_photo(path):
    try:
        img = Image.open(path)
        w, h = img.size
        if w < MIN_PHOTO_SIZE or h < MIN_PHOTO_SIZE:
            return False
        ratio = w / h
        if ratio < MIN_ASPECT_RATIO or ratio > MAX_ASPECT_RATIO:
            return False
        img = img.convert("RGB")
        if get_saturation_mean(img) < MIN_SATURATION:
            return False
        if dominant_background_ratio(img) > MAX_BACKGROUND_RATIO:
            return False
        if ImageStat.Stat(img.convert("L")).stddev[0] < MIN_STD_DEV:
            return False
        return True
    except Exception:
        return False


def make_thumbnail(src_path, dest_path):
    img = Image.open(src_path).convert("RGB")
    img.thumbnail((THUMB_MAX_W, THUMB_MAX_H), Image.Resampling.LANCZOS)
    img.save(dest_path, "JPEG", quality=85)


def process_all_documents(input_dir: Path, temp_folder: Path, progress=None, warnings=None):
    warnings = warnings if warnings is not None else []
    temp_folder.mkdir(parents=True, exist_ok=True)

    input_files = collect_input_files(input_dir)
    total = len(input_files)
    raw_valid_photos = []

    for idx, (full_path, rel_name) in enumerate(input_files, start=1):
        if progress:
            progress(0.05 + 0.60 * (idx - 1) / max(total, 1), f"Membaca dokumen {idx}/{total}: {rel_name}")
        prefix = f"{idx:04d}_{slugify(rel_name)[:40]}"
        if full_path.lower().endswith(".pdf"):
            extracted = extract_photos_from_pdf(full_path, str(temp_folder), prefix, warnings)
        else:
            extracted = extract_photos_from_docx(full_path, str(temp_folder), prefix, warnings)

        for p, label in extracted:
            if is_real_photo(p):
                raw_valid_photos.append({"doc_name": rel_name, "file_name": label, "full_temp_path": p})

    if progress:
        progress(0.68, "Menghitung kemiripan foto...")

    hashes = {}
    doc_appearance = defaultdict(set)
    for item in raw_valid_photos:
        try:
            h = imagehash.phash(Image.open(item["full_temp_path"]))
            hashes[item["full_temp_path"]] = h
            doc_appearance[str(h)].add(item["doc_name"])
        except Exception:
            hashes[item["full_temp_path"]] = None

    template_hashes = {h for h, docs in doc_appearance.items() if len(docs) > MAX_DOC_OCCURRENCES}

    all_valid_photos = []
    for item in raw_valid_photos:
        h = hashes[item["full_temp_path"]]
        if h is not None and str(h) in template_hashes:
            continue
        all_valid_photos.append(item)

    duplicate_groups = []
    visited = set()
    paths = [it["full_temp_path"] for it in all_valid_photos if hashes[it["full_temp_path"]] is not None]

    for i in range(len(paths)):
        p1 = paths[i]
        if p1 in visited:
            continue
        group = [p1]
        for j in range(i + 1, len(paths)):
            p2 = paths[j]
            if p2 in visited:
                continue
            if hashes[p1] - hashes[p2] <= HASH_DISTANCE_THRESHOLD:
                group.append(p2)
                visited.add(p2)
        if len(group) > 1:
            visited.add(p1)
            duplicate_groups.append(group)

    return input_files, all_valid_photos, duplicate_groups


def build_excel_report(all_photos, duplicate_groups, output_file: Path, thumb_folder: Path):
    thumb_folder.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws_dup = wb.active
    ws_dup.title = "Grup Foto Duplikat"
    headers_dup = ["Grup ID", "Dokumen BAST (Foto 1)", "File Foto 1", "Screenshot Foto 1",
                   "Dokumen BAST Kembaran (Foto 2)", "File Foto 2", "Screenshot Foto 2"]
    ws_dup.append(headers_dup)
    for col, w in zip("ABCDEFG", [10, 45, 22, COL_WIDTH_SCREENSHOT, 45, 22, COL_WIDTH_SCREENSHOT]):
        ws_dup.column_dimensions[col].width = w

    path_to_info = {p["full_temp_path"]: p for p in all_photos}

    thumb_idx = 1
    for g_id, group in enumerate(duplicate_groups, start=1):
        p1 = group[0]
        info1 = path_to_info[p1]
        for p2 in group[1:]:
            info2 = path_to_info[p2]
            row_idx = ws_dup.max_row + 1
            ws_dup.cell(row=row_idx, column=1, value=f"Grup #{g_id}")
            ws_dup.cell(row=row_idx, column=2, value=info1["doc_name"])
            ws_dup.cell(row=row_idx, column=3, value=info1["file_name"])
            ws_dup.cell(row=row_idx, column=5, value=info2["doc_name"])
            ws_dup.cell(row=row_idx, column=6, value=info2["file_name"])
            ws_dup.row_dimensions[row_idx].height = ROW_HEIGHT_PT

            t1 = thumb_folder / f"thumb_{thumb_idx}_1.jpg"
            make_thumbnail(p1, t1)
            ws_dup.add_image(XLImage(str(t1)), f"D{row_idx}")
            t2 = thumb_folder / f"thumb_{thumb_idx}_2.jpg"
            make_thumbnail(p2, t2)
            ws_dup.add_image(XLImage(str(t2)), f"G{row_idx}")
            thumb_idx += 1

    for col in range(1, 8):
        c = ws_dup.cell(row=1, column=col)
        c.fill, c.font = FILL_HEADER, FONT_HEADER
        c.alignment = Alignment(horizontal="center", vertical="center")
    for r in range(2, ws_dup.max_row + 1):
        for c in range(1, 8):
            cell = ws_dup.cell(row=r, column=c)
            cell.fill = FILL_DUPLIKAT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    ws_all = wb.create_sheet(title="Semua Foto")
    ws_all.append(["No", "Nama Dokumen BAST", "Nama File Foto", "Status", "Screenshot"])
    for col, w in zip("ABCDE", [5, 45, 22, 15, COL_WIDTH_SCREENSHOT]):
        ws_all.column_dimensions[col].width = w

    dup_paths_set = {p for grp in duplicate_groups for p in grp}
    for idx, item in enumerate(all_photos, start=1):
        row_idx = idx + 1
        is_dup = item["full_temp_path"] in dup_paths_set
        ws_all.cell(row=row_idx, column=1, value=idx)
        ws_all.cell(row=row_idx, column=2, value=item["doc_name"])
        ws_all.cell(row=row_idx, column=3, value=item["file_name"])
        s_cell = ws_all.cell(row=row_idx, column=4, value="DUPLIKAT" if is_dup else "OK")
        ws_all.row_dimensions[row_idx].height = ROW_HEIGHT_PT

        t = thumb_folder / f"thumb_all_{idx}.jpg"
        make_thumbnail(item["full_temp_path"], t)
        ws_all.add_image(XLImage(str(t)), f"E{row_idx}")

        for col in range(1, 6):
            c = ws_all.cell(row=row_idx, column=col)
            c.fill = FILL_DUPLIKAT if is_dup else FILL_OK
            c.border = THIN_BORDER
            c.alignment = Alignment(vertical="center", wrap_text=True)
        s_cell.font = FONT_DUPLIKAT if is_dup else FONT_OK

    for col in range(1, 6):
        c = ws_all.cell(row=1, column=col)
        c.fill, c.font = FILL_HEADER, FONT_HEADER
        c.alignment = Alignment(horizontal="center", vertical="center")

    wb.save(str(output_file))


def run_bast_check(input_dir: Path, work_dir: Path, progress=None, max_preview_groups=20):
    """
    Jalankan seluruh pipeline. Return dict:
      excel_bytes, dup_zip_bytes, n_docs, n_photos, n_groups, n_dup_photos,
      preview (list grup -> list (doc_name, file_name, jpeg_bytes)), warnings
    """
    warnings = []
    temp_folder = work_dir / "_temp_extracted_photos"
    thumb_folder = work_dir / "_temp_thumbnails"
    out_xlsx = work_dir / "hasil_laporan_semua_bast.xlsx"

    for d in (temp_folder, thumb_folder):
        if d.exists():
            shutil.rmtree(d)

    input_files, all_photos, groups = process_all_documents(input_dir, temp_folder, progress, warnings)
    if not input_files:
        raise ValueError("Tidak ada file PDF/DOCX yang ditemukan di file yang diupload.")

    if progress:
        progress(0.80, "Membuat laporan Excel...")
    build_excel_report(all_photos, groups, out_xlsx, thumb_folder)
    excel_bytes = out_xlsx.read_bytes()

    if progress:
        progress(0.92, "Menyiapkan foto duplikat...")
    path_to_info = {p["full_temp_path"]: p for p in all_photos}

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for g_id, group in enumerate(groups, start=1):
            for p in group:
                zf.write(p, arcname=f"Grup_{g_id:03d}/{os.path.basename(p)}")

    preview = []
    for g_id, group in enumerate(groups[:max_preview_groups], start=1):
        items = []
        for p in group[:4]:
            im = Image.open(p).convert("RGB")
            im.thumbnail((320, 240))
            b = io.BytesIO()
            im.save(b, "JPEG", quality=80)
            info = path_to_info[p]
            items.append((info["doc_name"], info["file_name"], b.getvalue()))
        preview.append({"group_id": g_id, "size": len(group), "items": items})

    for d in (temp_folder, thumb_folder):
        shutil.rmtree(d, ignore_errors=True)

    if progress:
        progress(1.0, "Selesai")

    return {
        "excel_bytes": excel_bytes,
        "dup_zip_bytes": buf.getvalue(),
        "n_docs": len(input_files),
        "n_photos": len(all_photos),
        "n_groups": len(groups),
        "n_dup_photos": sum(len(g) for g in groups),
        "preview": preview,
        "warnings": warnings,
    }
